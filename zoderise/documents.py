from __future__ import annotations

import io
import os
import re
import secrets
import time
import zipfile
import unicodedata
from collections import Counter
from pathlib import Path
from typing import Any, Optional, Sequence

from PIL import Image, ImageOps, UnidentifiedImageError

from .config import settings
from .database import (
    audit,
    find_duplicate_document,
    get_document_for_access,
    get_retrieval_rows,
    insert_chunks,
    insert_document_record,
    update_document_status,
)
from .security import secure_filename_stem, sha256_bytes

try:
    import docx
except Exception:
    docx = None

try:
    import pymupdf as fitz
except Exception:
    try:
        import fitz  # type: ignore
    except Exception:
        fitz = None

try:
    import PyPDF2
except Exception:
    PyPDF2 = None

try:
    import pytesseract
    pytesseract.get_tesseract_version()
    OCR_AVAILABLE = True
except Exception:
    pytesseract = None
    OCR_AVAILABLE = False


ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md", ".png", ".jpg", ".jpeg", ".webp"}
DOCUMENT_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp"}

UPLOAD_ROOT = settings.storage_root / "uploads"
UPLOAD_ROOT.mkdir(parents=True, exist_ok=True)


def normalize_text(text: str) -> str:
    text = str(text or "").replace("\x00", "")
    text = re.sub(r"\r\n?", "\n", text)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def available_ocr_languages() -> str:
    if not OCR_AVAILABLE or pytesseract is None:
        return "eng"
    try:
        installed = set(pytesseract.get_languages(config=""))
        requested = [x for x in settings.ocr_lang.split("+") if x in installed]
        if requested:
            return "+".join(requested)
        return "eng" if "eng" in installed else (next(iter(installed)) if installed else "eng")
    except Exception:
        return "eng"


def ocr_image(image: Image.Image) -> str:
    if not OCR_AVAILABLE or pytesseract is None:
        return ""
    return normalize_text(pytesseract.image_to_string(image, lang=available_ocr_languages()))


def chunk_text(text: str, max_chars: int | None = None, overlap: int | None = None) -> list[str]:
    text = normalize_text(text)
    if not text:
        return []
    max_chars = max_chars or settings.rag_chunk_size
    overlap = settings.rag_chunk_overlap if overlap is None else overlap
    overlap = min(overlap, max_chars // 3)
    output: list[str] = []
    start = 0
    n = len(text)
    while start < n:
        end = min(start + max_chars, n)
        if end < n:
            floor = start + int(max_chars * 0.55)
            candidates = [
                text.rfind("\n\n", floor, end),
                text.rfind(". ", floor, end),
                text.rfind("؟ ", floor, end),
                text.rfind("! ", floor, end),
                text.rfind("; ", floor, end),
            ]
            boundary = max(candidates)
            if boundary > start:
                end = boundary + 1
        part = text[start:end].strip()
        if part:
            output.append(part)
        if end >= n:
            break
        start = max(start + 1, end - overlap)
    return output


def _validate_docx_archive(raw: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as zf:
            infos = zf.infolist()
            total = sum(max(0, int(i.file_size)) for i in infos)
            if total > settings.max_docx_uncompressed_mb * 1024 * 1024:
                raise ValueError("DOCX expands beyond the configured safe uncompressed-size limit.")
            if len(infos) > 10_000:
                raise ValueError("DOCX contains an excessive number of archive entries.")
            for info in infos:
                normalized = Path(info.filename)
                if normalized.is_absolute() or ".." in normalized.parts:
                    raise ValueError("DOCX contains an unsafe archive path.")
    except zipfile.BadZipFile as exc:
        raise ValueError("The DOCX archive is invalid or corrupted.") from exc


def _validate_image(raw: bytes) -> None:
    try:
        with Image.open(io.BytesIO(raw)) as image:
            width, height = image.size
            if width <= 0 or height <= 0 or width * height > settings.max_image_pixels:
                raise ValueError("Image dimensions exceed the configured safety limit.")
            image.verify()
    except (UnidentifiedImageError, OSError) as exc:
        raise ValueError("The image is invalid or corrupted.") from exc


def validate_upload(filename: str, raw: bytes) -> str:
    if not filename or len(filename) > 180:
        raise ValueError("Invalid filename.")
    ext = Path(filename).suffix.lower()
    if ext not in ALLOWED_EXTENSIONS:
        raise ValueError(f"Unsupported file type: {ext or 'unknown'}")
    if not raw:
        raise ValueError("The uploaded file is empty.")
    if len(raw) > settings.max_upload_mb * 1024 * 1024:
        raise ValueError(f"File exceeds the {settings.max_upload_mb} MB per-file limit.")
    signatures = {
        ".pdf": raw.startswith(b"%PDF-"),
        ".png": raw.startswith(b"\x89PNG\r\n\x1a\n"),
        ".jpg": raw.startswith(b"\xff\xd8\xff"),
        ".jpeg": raw.startswith(b"\xff\xd8\xff"),
        ".webp": len(raw) > 12 and raw[:4] == b"RIFF" and raw[8:12] == b"WEBP",
        ".docx": raw.startswith(b"PK\x03\x04"),
    }
    if ext in signatures and not signatures[ext]:
        raise ValueError("File contents do not match the filename extension.")
    if ext == ".docx":
        _validate_docx_archive(raw)
    if ext in IMAGE_EXTENSIONS:
        _validate_image(raw)
    return ext


def _extract_pdf(raw: bytes) -> tuple[str, list[dict[str, Any]]]:
    if fitz is not None:
        pdf = fitz.open(stream=raw, filetype="pdf")
        try:
            if pdf.page_count > settings.max_pdf_pages:
                raise ValueError(f"PDF exceeds the {settings.max_pdf_pages}-page limit.")
            pages: list[dict[str, Any]] = []
            all_text: list[str] = []
            for page_number, page in enumerate(pdf, 1):
                text = normalize_text(page.get_text("text"))
                if not text and OCR_AVAILABLE:
                    area = max(1.0, page.rect.width * page.rect.height)
                    scale = min(1.6, (settings.max_image_pixels / area) ** 0.5 * 0.99)
                    pix = page.get_pixmap(matrix=fitz.Matrix(scale, scale), alpha=False)
                    if pix.width * pix.height <= settings.max_image_pixels:
                        image = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
                        text = ocr_image(image)
                pages.append({"page": page_number, "text": text})
                if text:
                    all_text.append(text)
            return "\n\n".join(all_text), pages
        finally:
            pdf.close()
    if PyPDF2 is not None:
        reader = PyPDF2.PdfReader(io.BytesIO(raw))
        if len(reader.pages) > settings.max_pdf_pages:
            raise ValueError(f"PDF exceeds the {settings.max_pdf_pages}-page limit.")
        pages = []
        all_text = []
        for page_number, page in enumerate(reader.pages, 1):
            text = normalize_text(page.extract_text() or "")
            pages.append({"page": page_number, "text": text})
            if text:
                all_text.append(text)
        return "\n\n".join(all_text), pages
    raise RuntimeError("No PDF parser is installed. Install PyMuPDF or PyPDF2.")


def _extract_docx(raw: bytes) -> tuple[str, list[dict[str, Any]]]:
    if docx is None:
        raise RuntimeError("python-docx is not installed.")
    document = docx.Document(io.BytesIO(raw))
    blocks: list[str] = []
    for paragraph in document.paragraphs:
        text = normalize_text(paragraph.text)
        if text:
            blocks.append(text)
    for table in document.tables:
        for row in table.rows:
            cells = [normalize_text(cell.text) for cell in row.cells]
            if any(cells):
                blocks.append(" | ".join(cells))
    text = "\n".join(blocks)
    return text, [{"page": None, "text": text}]


def extract_file(filename: str, raw: bytes) -> tuple[str, list[dict[str, Any]]]:
    ext = validate_upload(filename, raw)
    if ext in {".txt", ".md"}:
        text = normalize_text(raw.decode("utf-8", errors="replace"))
        return text, [{"page": None, "text": text}]
    if ext == ".pdf":
        return _extract_pdf(raw)
    if ext == ".docx":
        return _extract_docx(raw)
    if ext in IMAGE_EXTENSIONS:
        with Image.open(io.BytesIO(raw)) as image:
            text = ocr_image(image.convert("RGB"))
        return text, [{"page": None, "text": text}]
    raise ValueError("Unsupported file type.")


def safe_storage_path(user_id: int, filename: str) -> Path:
    ext = Path(filename).suffix.lower()
    if ext not in ALLOWED_EXTENSIONS:
        raise ValueError("Unsupported file type.")
    user_dir = UPLOAD_ROOT / f"user_{int(user_id)}"
    user_dir.mkdir(parents=True, exist_ok=True)
    return user_dir / f"{int(time.time())}_{secrets.token_hex(8)}_{secure_filename_stem(filename)}{ext}"


def _safe_remove_file(path_value: str | Path, user_id: Optional[int] = None) -> bool:
    path = Path(path_value)
    root = UPLOAD_ROOT.resolve()
    try:
        resolved = path.resolve()
        resolved.relative_to(root)
        if user_id is not None:
            resolved.relative_to((UPLOAD_ROOT / f"user_{int(user_id)}").resolve())
    except Exception:
        return False
    try:
        if resolved.is_file():
            resolved.unlink()
            return True
    except OSError:
        pass
    return False


def save_chat_image(user_id: int, filename: str, raw: bytes) -> dict[str, Any]:
    ext = validate_upload(filename, raw)
    if ext not in IMAGE_EXTENSIONS:
        raise ValueError("Only image attachments can be stored as chat images.")
    with Image.open(io.BytesIO(raw)) as image:
        image = ImageOps.exif_transpose(image).convert("RGB")
        image.thumbnail((2560, 2560))
        if image.width * image.height > settings.max_image_pixels:
            raise ValueError("Image is too large.")
        buf = io.BytesIO()
        image.save(buf, format="JPEG", quality=88, optimize=True)
    image_bytes = buf.getvalue()
    path = safe_storage_path(user_id, f"{Path(filename).stem}.jpg")
    path.write_bytes(image_bytes)
    return {
        "filename": Path(filename).name[:180],
        "stored_path": str(path),
        "stored_format": "JPEG",
        "size_bytes": len(image_bytes),
    }


def remove_chat_attachments(attachments: Sequence[dict[str, Any]], user_id: int) -> int:
    removed = 0
    for item in attachments:
        path = item.get("stored_path") if isinstance(item, dict) else None
        if path and _safe_remove_file(path, user_id=user_id):
            removed += 1
    return removed


def _build_chunk_records(pages: Sequence[dict[str, Any]]) -> list[tuple[Optional[int], str, str]]:
    chunks: list[tuple[Optional[int], str, str]] = []
    for page in pages:
        page_num = page.get("page")
        for part in chunk_text(page.get("text", "")):
            chunks.append((page_num, "", part))
    return chunks


def save_document(
    user_id: int,
    filename: str,
    raw: bytes,
    mime_type: str,
    *,
    domain: str,
    visibility: str = "private",
    tags: Sequence[str] = (),
) -> tuple[int, int, str]:
    """Extract and index a local document without requiring a third AI model.

    HORIZON V8 intentionally uses lexical/BM25-style retrieval so the complete
    platform only depends on `zoderise:latest` and `zoderise-vision:latest`.
    """
    ext = validate_upload(filename, raw)
    if visibility not in {"private", "shared"}:
        visibility = "private"
    checksum = sha256_bytes(raw)
    existing = find_duplicate_document(user_id, checksum)
    if existing:
        return int(existing["id"]), 0, "duplicate"

    text, pages = extract_file(filename, raw)
    if not text:
        raise ValueError("No readable text was extracted. Scanned PDFs and image indexing need OCR. Use Chat to analyze images with the vision model.")

    path = safe_storage_path(user_id, filename)
    path.write_bytes(raw)
    doc_id: Optional[int] = None
    try:
        doc_id = insert_document_record(
            user_id,
            Path(filename).name,
            str(path),
            mime_type or "application/octet-stream",
            len(raw),
            checksum,
            text,
            "indexing",
            domain,
            visibility,
            tags,
        )
        raw_chunks = _build_chunk_records(pages)
        records = [(page, section, part, None) for page, section, part in raw_chunks]
        if records:
            insert_chunks(user_id, doc_id, records)
        update_document_status(doc_id, "ready")
        audit(user_id, "document_indexed", {"document_id": doc_id, "chunks": len(records), "visibility": visibility})
        return doc_id, len(records), "ready"
    except Exception:
        _safe_remove_file(path, user_id=user_id)
        if doc_id is not None:
            try:
                from .database import connect
                with connect() as conn:
                    conn.execute("DELETE FROM documents WHERE id=? AND user_id=?", (doc_id, user_id))
                    conn.commit()
            except Exception:
                pass
        raise


def process_ephemeral_document(filename: str, raw: bytes) -> list[dict[str, Any]]:
    text, pages = extract_file(filename, raw)
    if not text:
        return []
    chunks = _build_chunk_records(pages)
    return [
        {"filename": Path(filename).name[:180], "page": page, "section": section, "text": part, "score": 1.0}
        for page, section, part in chunks
    ]


def _query_terms(text: str) -> list[str]:
    text = unicodedata.normalize("NFKC", text).translate(str.maketrans({"ي": "ی", "ك": "ک", "\u200c": " ", "ـ": ""}))
    text = "".join(c for c in text if not unicodedata.combining(c))
    stopwords = {"the", "a", "an", "is", "are", "of", "to", "and", "in", "for", "what", "how", "this", "that"}
    return [t.lower() for t in re.findall(r"[\w]+", text, flags=re.UNICODE) if t.lower() not in stopwords]


def _lexical_rank(query: str, texts: Sequence[str]) -> list[float]:
    """Small BM25-inspired local scorer with phrase and coverage bonuses.

    It works for English and Unicode word characters (including Dari/Pashto)
    and requires no embedding model or network access.
    """
    terms = _query_terms(query)
    if not terms or not texts:
        return [0.0] * len(texts)
    unique_terms = list(dict.fromkeys(terms))
    tokenized = [_query_terms(text) for text in texts]
    lengths = [max(1, len(tokens)) for tokens in tokenized]
    avg_len = sum(lengths) / max(1, len(lengths))
    n_docs = len(texts)
    doc_freq = Counter(term for tokens in tokenized for term in set(tokens))
    k1, b = 1.35, 0.72
    phrase = query.strip().lower()
    scores: list[float] = []
    for raw_text, tokens, length in zip(texts, tokenized, lengths):
        counts: dict[str, int] = {}
        for token in tokens:
            counts[token] = counts.get(token, 0) + 1
        score = 0.0
        matched = 0
        for term in unique_terms:
            tf = counts.get(term, 0)
            if not tf:
                continue
            matched += 1
            df = doc_freq.get(term, 0)
            idf = max(0.05, __import__("math").log(1.0 + (n_docs - df + 0.5) / (df + 0.5)))
            denom = tf + k1 * (1 - b + b * length / max(1.0, avg_len))
            score += idf * ((tf * (k1 + 1)) / denom)
        coverage = matched / max(1, len(unique_terms))
        score *= 0.72 + 0.55 * coverage
        low = raw_text.lower()
        if len(phrase) >= 4 and phrase in low:
            score += 1.8
        if query.strip() and low.startswith(query.strip().lower()):
            score += 0.35
        scores.append(score)
    max_score = max(scores, default=0.0)
    if max_score <= 0:
        return [0.0] * len(scores)
    return [min(1.0, score / max_score) for score in scores]


def lexical_score(query: str, text: str) -> float:
    return _lexical_rank(query, [text])[0] if text else 0.0


def retrieve_chunks(
    user_id: int,
    query: str,
    *,
    top_k: Optional[int] = None,
    document_ids: Optional[Sequence[int]] = None,
) -> list[dict[str, Any]]:
    rows = get_retrieval_rows(user_id, document_ids)
    if not rows:
        return []
    top_k = top_k or settings.rag_top_k
    texts = [str(row["text"] or "") for row in rows]
    scores = _lexical_rank(query, texts)
    ranked = sorted(zip(scores, rows), key=lambda item: item[0], reverse=True)
    results: list[dict[str, Any]] = []
    for score, row in ranked:
        if len(results) >= top_k:
            break
        if score <= 0:
            continue
        if score < settings.rag_similarity_threshold:
            continue
        results.append({
            "score": round(float(score), 4),
            "filename": row["filename"],
            "page": row["page"],
            "section": row["section"],
            "text": row["text"],
            "chunk_id": int(row["id"]),
            "document_id": int(row["document_id"]),
            "visibility": row["visibility"],
            "domain": row["domain"],
        })
    return results


def retrieve_ephemeral(query: str, chunks: Sequence[dict[str, Any]], top_k: Optional[int] = None) -> list[dict[str, Any]]:
    top_k = top_k or settings.rag_top_k
    if not chunks:
        return []
    texts = [str(chunk.get("text", "")) for chunk in chunks]
    scores = _lexical_rank(query, texts)
    ranked: list[tuple[float, dict[str, Any]]] = []
    for score, chunk in zip(scores, chunks):
        item = dict(chunk)
        item["score"] = round(float(score), 4)
        ranked.append((score, item))
    ranked.sort(key=lambda item: item[0], reverse=True)
    return [item for score, item in ranked[:top_k] if score > 0 and score >= settings.rag_similarity_threshold]


def delete_stored_document(user_id: int, document_id: int, is_admin: bool = False) -> bool:
    from .database import delete_document

    path = delete_document(user_id, document_id, is_admin=is_admin)
    if not path:
        return False
    _safe_remove_file(path)
    audit(user_id, "document_deleted", {"document_id": document_id})
    return True


def document_preview(user_id: int, document_id: int, max_chars: int = 5000) -> Optional[dict[str, Any]]:
    row = get_document_for_access(user_id, document_id)
    if not row:
        return None
    return {
        "id": int(row["id"]),
        "filename": row["filename"],
        "domain": row["domain"],
        "visibility": row["visibility"],
        "owner_name": row["owner_name"],
        "text": str(row["extracted_text"] or "")[:max_chars],
        "status": row["status"],
        "size_bytes": int(row["size_bytes"]),
        "created_at": float(row["created_at"]),
    }


def file_capabilities() -> dict[str, Any]:
    return {
        "ocr_available": OCR_AVAILABLE,
        "ocr_languages": available_ocr_languages() if OCR_AVAILABLE else "",
        "pdf_parser": "PyMuPDF" if fitz is not None else ("PyPDF2" if PyPDF2 is not None else "unavailable"),
        "docx_parser": bool(docx is not None),
        "allowed_extensions": sorted(ALLOWED_EXTENSIONS),
    }


def remove_user_storage(user_id: int) -> int:
    user_dir = UPLOAD_ROOT / f"user_{int(user_id)}"
    if not user_dir.exists():
        return 0
    removed = 0
    for path in sorted(user_dir.rglob("*"), reverse=True):
        try:
            if path.is_file():
                path.unlink()
                removed += 1
            elif path.is_dir():
                path.rmdir()
        except OSError:
            pass
    try:
        user_dir.rmdir()
    except OSError:
        pass
    return removed
